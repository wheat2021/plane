#!/usr/bin/env python3
"""
生成历史迭代需求 import_ready CSV（0131/0307/0328）
"""
import csv
import json
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

# Patch openpyxl DataValidation bug
_orig_init = DataValidation.__init__
def _patched_init(self, **kwargs):
    kwargs.pop("id", None)
    _orig_init(self, **kwargs)
DataValidation.__init__ = _patched_init

import warnings
warnings.filterwarnings("ignore")

# ── 常量配置 ──────────────────────────────────────────────────────────────────
PROJECT_ID  = "18b7ccc8-4b96-4af0-8c6f-76550cba28a7"
TYPE_ID     = "5e8c9352-8553-4e87-97dc-88594010f83a"   # Requirement
STATE_ID    = "4657dc96-1572-41e3-becd-0ba3d4ff8e22"   # Backlog (default)

CYCLE_IDS = {
    "0131": "b5ee98b6-f0a4-4371-8999-c080a6a5b3df",
    "0307": "eb4e80fd-280b-4ed8-9213-b04ea8f09627",
    "0328": "8b6f3266-ad5d-4193-8b2e-724da8b4a476",
}

MODULE_IDS = {
    "FICC策略平台整合项目":                     "d3aec2ab-36f3-4308-a0fb-83e8e501a6a4",
    "O45迁移项目":                              "7218809e-4f32-4466-a207-8260f28b501d",
    "【Calypso下线项目】Tarf簿记":              "1b35a6b5-f222-4016-9fb1-ad7d924bded4",
    "【Calypso下线项目】美国国债簿记":          "cf140190-a666-41a7-9ee5-5432f20d38c3",
    "【海外交易能力建设】香港外汇期货做市项目": "770f76de-e4c8-4d07-a092-9639b2a1ab52",
    "【海外交易能力建设】美国国债自营交易链路": "68557d3f-59d5-48dd-b6c9-cdecd7adffdf",
    "FICC人工智能专项项目":                     "f7740574-83da-4168-aea4-3bed1c0cf870",
    "FICC交易日历优化项目":                     "c05328ce-9a9c-479f-8b99-5a03c005794a",
    "境外本地化部署规划":                       "3a12f818-b40c-4a20-a18f-d55b6256056b",
    "FICCUI测试专项":                           "9687131e-1b1b-4447-815e-6dc75b38be73",
}

IT_PM_UUIDS = {
    "夏清":   "9a969130-2666-4604-9326-c2b17ca09ef1",
    "刘佩金": "daa28f62-9d59-4c8a-9c93-28d931e72b99",
    "吴克易琼": "943c1313-5035-46a0-b1e2-91e3a19546a6",
    "曾慧聪": "86316c00-0847-43a8-8176-834f264e0307",
    "成楠":   "6fe8459b-3b6b-4197-a29a-f95f2a7e20cc",
    "钱高翔": "535b3838-1804-4a1f-8353-5320babd97b8",
    "姜玉珠": "15e4d327-6119-4941-85ac-624f7b021630",
    "孙玉涵": "8baed4ad-64c0-4922-b4bb-e9d77a865fd7",
    "毕成功": "b5cf400d-4a1a-4daf-8ad7-dc33d720391c",
    "陈宇":   "5ac3485a-8d93-454c-bd0d-0c048df1ba42",
    "陈治平": "77342d2f-a6cf-4a2d-8142-2f5f84d0636a",
    "许可":   "14b265fe-200d-46aa-8e25-66f57a1e130d",
    "胡于凡": "1f032804-4e24-4b85-a5d3-e57bb39f84b4",
    "朱泓飞": "b0c16d7f-e6db-46ed-b34e-ff2059b1c211",
}

CSV_COLUMNS = [
    "external_id", "name", "description_html", "type_id", "state_id",
    "external_source", "assignees", "extra_properties", "cycle_id", "module_id"
]

def cell(row, idx):
    """1-indexed, returns stripped string or None"""
    if idx - 1 >= len(row):
        return None
    v = row[idx - 1]
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None

def resolve_it_pm(raw):
    """处理复合值（钱高翔/朱泓飞 取后者），返回 UUID 或 None"""
    if not raw:
        return None
    # 复合值：取最后一个
    parts = [p.strip() for p in raw.replace("，", "/").split("/") if p.strip()]
    name = parts[-1] if parts else raw
    uuid = IT_PM_UUIDS.get(name)
    if not uuid:
        print(f"  ⚠ IT_PM 未匹配: {raw!r}")
    return uuid

def make_description_html(text):
    if not text:
        return ""
    # 简单段落 HTML
    paras = str(text).strip().split("\n")
    return "".join(f"<p>{p.strip()}</p>" for p in paras if p.strip())

def read_sheet(path, sheet_name):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet_name]
    rows = []
    for r in range(2, ws.max_row + 1):
        row = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if all(v is None for v in row):
            continue
        rows.append(row)
    return rows

def write_csv(output_path, records):
    with open(output_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=CSV_COLUMNS)
        writer.writeheader()
        writer.writerows(records)
    print(f"  写入 {len(records)} 行 → {output_path}")

# ── 0131 ──────────────────────────────────────────────────────────────────────
# 列映射（1-indexed）：
#   col1=所属部门 col2=需求编号 col3=概要(name) col4=境内/外
#   col5=一级分类 col6=二级分类 col7=三级分类 col8=所属中心 col9=所属团队
#   col10=需求提出人 col11=固收PM col12=IT产品经理 col13=IT技术负责人
#   col14=需求是否准入 col15=是否纳入交付 col16=预估排期 col17=当前迭代交付内容
def process_0131():
    print("处理 0131...")
    rows = read_sheet("26-0131迭代需求管理（业务+内生）.xlsx", "0131-需求列表")
    records = []
    skip = 0
    for row in rows:
        name_raw = cell(row, 3)
        if not name_raw:
            skip += 1
            continue
        name = f"{name_raw} [0131]"
        ext_id = cell(row, 2) or ""
        delivery = cell(row, 17)
        desc_html = make_description_html(delivery)
        it_pm_raw = cell(row, 12)
        in_delivery_raw = cell(row, 15)
        estimated = cell(row, 16)

        ep = {
            "department":         cell(row, 1),
            "domestic_overseas":  cell(row, 4),
            "category_l1":        cell(row, 5),
            "category_l2":        cell(row, 6),
            "category_l3":        cell(row, 7),
            "center":             cell(row, 8),
            "team":               cell(row, 9),
            "biz_pm":             cell(row, 11),
            "it_pm":              resolve_it_pm(it_pm_raw),
            "techLead":           cell(row, 13),
            "in_delivery":        True if in_delivery_raw == "是" else False,
            "estimated_iteration": estimated,
            "delivery_content":   delivery,
        }
        # 清除 None 值
        ep = {k: v for k, v in ep.items() if v is not None and v != False or k == "in_delivery"}

        records.append({
            "external_id":    ext_id,
            "name":           name[:255],
            "description_html": desc_html,
            "type_id":        TYPE_ID,
            "state_id":       STATE_ID,
            "external_source": "jira",
            "assignees":      "[]",
            "extra_properties": json.dumps(ep, ensure_ascii=False),
            "cycle_id":       CYCLE_IDS["0131"],
            "module_id":      "",
        })
    print(f"  跳过无名行: {skip}")
    write_csv("26-0131_import_ready.csv", records)
    return records

# ── 0307 ──────────────────────────────────────────────────────────────────────
# col1=所属部门 col2=需求编号 col3=需求名 col4=需求描述
# col5=境内/外 col6=一级分类 col7=二级分类 col8=三级分类
# col9=所属中心 col10=所属团队 col11=需求提出人 col12=固收PM
# col13=优先级 col14=IT产品经理 col15=IT技术负责人
# col16=需求是否准入 col17=是否纳入交付 col18=预估迭代
# col19=当前迭代交付内容 col20=备注
def process_0307():
    print("处理 0307...")
    rows = read_sheet("26-0307迭代需求(业务+内生).xlsx", "0307-需求列表")
    records = []
    skip = 0
    for row in rows:
        name_raw = cell(row, 3)
        if not name_raw:
            skip += 1
            continue
        name = f"{name_raw} [0307]"
        ext_id = cell(row, 2) or ""
        req_desc = cell(row, 4)
        delivery = cell(row, 19)
        # description_html: 需求描述（主要）+ 当前交付（补充）
        desc_parts = []
        if req_desc:
            desc_parts.append(req_desc)
        if delivery:
            desc_parts.append(f"【当前迭代交付】{delivery}")
        desc_html = make_description_html("\n".join(desc_parts))
        it_pm_raw = cell(row, 14)
        in_delivery_raw = cell(row, 17)
        estimated = cell(row, 18)

        ep = {
            "department":         cell(row, 1),
            "domestic_overseas":  cell(row, 5),
            "category_l1":        cell(row, 6),
            "category_l2":        cell(row, 7),
            "category_l3":        cell(row, 8),
            "center":             cell(row, 9),
            "team":               cell(row, 10),
            "biz_pm":             cell(row, 12),
            "biz_priority":       cell(row, 13),
            "it_pm":              resolve_it_pm(it_pm_raw),
            "techLead":           cell(row, 15),
            "in_delivery":        True if in_delivery_raw == "是" else False,
            "estimated_iteration": estimated,
            "delivery_content":   delivery,
            "remarks":            cell(row, 20),
        }
        ep = {k: v for k, v in ep.items() if v is not None and v != False or k == "in_delivery"}

        records.append({
            "external_id":    ext_id,
            "name":           name[:255],
            "description_html": desc_html,
            "type_id":        TYPE_ID,
            "state_id":       STATE_ID,
            "external_source": "jira",
            "assignees":      "[]",
            "extra_properties": json.dumps(ep, ensure_ascii=False),
            "cycle_id":       CYCLE_IDS["0307"],
            "module_id":      "",
        })
    print(f"  跳过无名行: {skip}")
    write_csv("26-0307_import_ready.csv", records)
    return records

# ── 0328 ──────────────────────────────────────────────────────────────────────
# col1=所属部门 col2=需求编号 col3=需求名 col4=需求描述
# col5=境内/外 col6=一级分类 col7=二级分类 col8=三级分类
# col9=所属中心 col10=所属团队 col11=需求提出人 col12=固收PM
# col13=优先级 col14=IT产品经理 col15=IT技术负责人
# col16=重点项目标签 col17=需求是否准入 col18=是否纳入交付
# col19=预估迭代 col20=预估迭代(副) col21=当前迭代交付内容 col22=备注
def process_0328():
    print("处理 0328...")
    rows = read_sheet("26-0328迭代需求(业务+内生).xlsx", "0328-需求列表")
    records = []
    skip = 0
    module_match_count = 0
    for row in rows:
        name_raw = cell(row, 3)
        if not name_raw:
            skip += 1
            continue
        name = f"{name_raw} [0328]"
        ext_id = cell(row, 2) or ""
        req_desc = cell(row, 4)
        delivery = cell(row, 21)
        desc_parts = []
        if req_desc:
            desc_parts.append(req_desc)
        if delivery:
            desc_parts.append(f"【当前迭代交付】{delivery}")
        desc_html = make_description_html("\n".join(desc_parts))
        it_pm_raw = cell(row, 14)
        key_project = cell(row, 16)
        in_delivery_raw = cell(row, 18)
        estimated = cell(row, 19)

        # Module 匹配
        module_id = ""
        if key_project and key_project != "无":
            module_id = MODULE_IDS.get(key_project, "")
            if module_id:
                module_match_count += 1
            else:
                print(f"  ⚠ 重点项目标签未匹配 Module: {key_project!r}")

        ep = {
            "department":         cell(row, 1),
            "domestic_overseas":  cell(row, 5),
            "category_l1":        cell(row, 6),
            "category_l2":        cell(row, 7),
            "category_l3":        cell(row, 8),
            "center":             cell(row, 9),
            "team":               cell(row, 10),
            "biz_pm":             cell(row, 12),
            "biz_priority":       cell(row, 13),
            "it_pm":              resolve_it_pm(it_pm_raw),
            "techLead":           cell(row, 15),
            "in_delivery":        True if in_delivery_raw == "是" else False,
            "estimated_iteration": estimated,
            "delivery_content":   delivery,
            "remarks":            cell(row, 22),
        }
        ep = {k: v for k, v in ep.items() if v is not None and v != False or k == "in_delivery"}

        records.append({
            "external_id":    ext_id,
            "name":           name[:255],
            "description_html": desc_html,
            "type_id":        TYPE_ID,
            "state_id":       STATE_ID,
            "external_source": "jira",
            "assignees":      "[]",
            "extra_properties": json.dumps(ep, ensure_ascii=False),
            "cycle_id":       CYCLE_IDS["0328"],
            "module_id":      module_id,
        })
    print(f"  跳过无名行: {skip}, Module 匹配: {module_match_count} 条")
    write_csv("26-0328_import_ready.csv", records)
    return records


if __name__ == "__main__":
    import os
    os.chdir(os.path.dirname(os.path.abspath(__file__)))
    r0131 = process_0131()
    r0307 = process_0307()
    r0328 = process_0328()
    print(f"\n总计: 0131={len(r0131)}, 0307={len(r0307)}, 0328={len(r0328)}, 合计={len(r0131)+len(r0307)+len(r0328)}")
